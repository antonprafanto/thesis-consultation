import docx
import sys

doc_skripsi = docx.Document('scratch/pedoman_docs/Template_Skripsi.docx')
doc_prop = docx.Document('scratch/pedoman_docs/Template_Proposal.docx')

def inspect_doc(name, doc):
    print(f"\n==========================================")
    print(f"=== DETAILED AUDIT: {name} ===")
    print(f"==========================================")
    
    # Sections & Margins
    for i, sec in enumerate(doc.sections):
        print(f"Section {i}: HeaderDist={sec.header_distance.cm if sec.header_distance else None}cm, FooterDist={sec.footer_distance.cm if sec.footer_distance else None}cm")
        print(f"  Top={sec.top_margin.cm:.2f}cm, Bottom={sec.bottom_margin.cm:.2f}cm, Left={sec.left_margin.cm:.2f}cm, Right={sec.right_margin.cm:.2f}cm")
        print(f"  DifferentFirstPage={sec.different_first_page_header_footer}")
        
        # Check header/footer content
        header_text = " ".join([p.text.strip() for p in sec.header.paragraphs if p.text.strip()])
        footer_text = " ".join([p.text.strip() for p in sec.footer.paragraphs if p.text.strip()])
        print(f"  Header: '{header_text}' | Footer: '{footer_text}'")

    # Tables
    print(f"\nTotal Tables: {len(doc.tables)}")
    for t_idx, t in enumerate(doc.tables):
        print(f"Table {t_idx} (rows={len(t.rows)}, cols={len(t.columns)}):")
        for r_idx in range(min(4, len(t.rows))):
            cells_text = [t.cell(r_idx, c).text.strip().replace('\n', ' ') for c in range(min(5, len(t.columns)))]
            print(f"  R{r_idx}: {' | '.join(cells_text)}")

    # Specific elements
    print("\n--- Key Paragraphs Inspection ---")
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if any(k in t.upper() for k in ['BAB ', 'LATAR BELAKANG', 'GAMBAR 2.1', 'TABEL 2.1', 'DAFTAR PUSTAKA', 'KATA PENGANTAR', 'PERNYATAAN KEASLIAN', 'HALAMAN PENGESAHAN', 'ABSTRAK', 'PERSAMAAN', '1.1']):
            runs_info = []
            for r in p.runs:
                font_name = r.font.name or "inherited"
                size = r.font.size.pt if r.font.size else "inherited"
                bold = "B" if r.font.bold else ""
                italic = "I" if r.font.italic else ""
                runs_info.append(f"{r.text.strip()[:20]}[{font_name},{size},{bold}{italic}]")
            align = str(p.alignment)
            sb = p.paragraph_format.space_before.pt if p.paragraph_format.space_before else 0
            sa = p.paragraph_format.space_after.pt if p.paragraph_format.space_after else 0
            ls = p.paragraph_format.line_spacing
            print(f"P{i:03d} (align={align}, sb={sb}pt, sa={sa}pt, ls={ls}): {t[:60]}")
            print(f"      Runs: {' | '.join(runs_info[:4])}")

inspect_doc("Template Skripsi", doc_skripsi)
inspect_doc("Template Proposal", doc_prop)
