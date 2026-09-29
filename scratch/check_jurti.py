import docx

doc = docx.Document('scratch/pedoman_docs/Template_JuRTI.docx')
print(f"JuRTI Paras: {len(doc.paragraphs)}, Tables: {len(doc.tables)}, Sections: {len(doc.sections)}")

for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt:
        if any(w in txt.lower() for w in ['abstrak', 'abstract', 'pendahuluan', 'metode', 'hasil', 'kesimpulan', 'daftar pustaka', 'referensi', 'penulis']):
            print(f"P{i:02d}: {txt[:100]}")
